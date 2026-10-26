"""Lê um CSV de vendas, limpa os dados e gera um relatório Excel formatado."""
import argparse
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUNAS_OBRIGATORIAS = {"data", "produto", "vendedor", "quantidade", "preco_unitario"}


def carregar_dados(caminho: str) -> pd.DataFrame:
    df = pd.read_csv(caminho)
    faltando = COLUNAS_OBRIGATORIAS - set(df.columns)
    if faltando:
        raise ValueError(f"Colunas ausentes no arquivo: {sorted(faltando)}")
    return df


def limpar_dados(df: pd.DataFrame) -> pd.DataFrame:
    """Remove duplicatas e linhas inválidas, padroniza textos e calcula o total."""
    df = df.copy()
    df["data"] = pd.to_datetime(df["data"], errors="coerce")
    df["quantidade"] = pd.to_numeric(df["quantidade"], errors="coerce")
    df["preco_unitario"] = pd.to_numeric(df["preco_unitario"], errors="coerce")
    df["vendedor"] = df["vendedor"].astype(str).str.strip().str.title()
    df["produto"] = df["produto"].astype(str).str.strip()

    df = df.dropna(subset=["data", "quantidade", "preco_unitario"])
    df = df[(df["quantidade"] > 0) & (df["preco_unitario"] > 0)]
    df = df.drop_duplicates()

    df["quantidade"] = df["quantidade"].astype(int)
    df["total"] = df["quantidade"] * df["preco_unitario"]
    return df.reset_index(drop=True)


def resumo_por_produto(df: pd.DataFrame) -> pd.DataFrame:
    return (
        df.groupby("produto", as_index=False)
        .agg(quantidade=("quantidade", "sum"), faturamento=("total", "sum"))
        .sort_values("faturamento", ascending=False)
        .reset_index(drop=True)
    )


def resumo_por_vendedor(df: pd.DataFrame) -> pd.DataFrame:
    return (
        df.groupby("vendedor", as_index=False)
        .agg(vendas=("total", "count"), faturamento=("total", "sum"))
        .sort_values("faturamento", ascending=False)
        .reset_index(drop=True)
    )


def resumo_mensal(df: pd.DataFrame) -> pd.DataFrame:
    mensal = df.assign(mes=df["data"].dt.to_period("M").astype(str))
    return (
        mensal.groupby("mes", as_index=False)
        .agg(faturamento=("total", "sum"))
        .sort_values("mes")
        .reset_index(drop=True)
    )


def _formatar_aba(ws, colunas_moeda: list[str]) -> None:
    cabecalho = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = cabecalho
        cell.alignment = Alignment(horizontal="center")
    nomes = [c.value for c in ws[1]]
    for idx, nome in enumerate(nomes, start=1):
        letra = get_column_letter(idx)
        largura = max(len(str(c.value)) if c.value is not None else 0 for c in ws[letra])
        ws.column_dimensions[letra].width = max(14, largura + 4)
        if nome in colunas_moeda:
            for cell in ws[letra][1:]:
                cell.number_format = 'R$ #,##0.00'
    ws.freeze_panes = "A2"


def _adicionar_grafico(ws, titulo: str, ancora: str) -> None:
    grafico = BarChart()
    grafico.title = titulo
    grafico.style = 10
    dados = Reference(ws, min_col=ws.max_column, min_row=1, max_row=ws.max_row)
    categorias = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
    grafico.add_data(dados, titles_from_data=True)
    grafico.set_categories(categorias)
    grafico.legend = None
    grafico.width, grafico.height = 18, 9
    ws.add_chart(grafico, ancora)


def gerar_relatorio(df: pd.DataFrame, saida: str) -> Path:
    destino = Path(saida)
    destino.parent.mkdir(parents=True, exist_ok=True)

    produtos = resumo_por_produto(df)
    vendedores = resumo_por_vendedor(df)
    mensal = resumo_mensal(df)
    kpis = pd.DataFrame(
        {
            "indicador": ["Faturamento total", "Nº de vendas", "Ticket médio"],
            "valor": [df["total"].sum(), len(df), df["total"].mean()],
        }
    )

    with pd.ExcelWriter(destino, engine="openpyxl") as writer:
        kpis.to_excel(writer, sheet_name="Resumo", index=False)
        produtos.to_excel(writer, sheet_name="Por Produto", index=False)
        vendedores.to_excel(writer, sheet_name="Por Vendedor", index=False)
        mensal.to_excel(writer, sheet_name="Mensal", index=False)
        df.assign(data=df["data"].dt.date).to_excel(writer, sheet_name="Dados Limpos", index=False)

    wb = load_workbook(destino)
    _formatar_aba(wb["Resumo"], [])
    for linha in wb["Resumo"].iter_rows(min_row=2):
        linha[1].number_format = 'R$ #,##0.00' if linha[0].value != "Nº de vendas" else "0"
    _formatar_aba(wb["Por Produto"], ["faturamento"])
    _formatar_aba(wb["Por Vendedor"], ["faturamento"])
    _formatar_aba(wb["Mensal"], ["faturamento"])
    _formatar_aba(wb["Dados Limpos"], ["preco_unitario", "total"])
    _adicionar_grafico(wb["Por Produto"], "Faturamento por produto", "E2")
    _adicionar_grafico(wb["Mensal"], "Faturamento mensal", "D2")
    wb.save(destino)
    return destino


def main() -> None:
    parser = argparse.ArgumentParser(description="Gera relatório Excel a partir de um CSV de vendas.")
    parser.add_argument("entrada", help="Caminho do CSV de vendas")
    parser.add_argument("-o", "--saida", default="output/relatorio_vendas.xlsx", help="Arquivo Excel de saída")
    args = parser.parse_args()

    bruto = carregar_dados(args.entrada)
    limpo = limpar_dados(bruto)
    print(f"Linhas lidas: {len(bruto)} | Linhas válidas: {len(limpo)} | Descartadas: {len(bruto) - len(limpo)}")
    caminho = gerar_relatorio(limpo, args.saida)
    print(f"Relatório gerado em: {caminho}")


if __name__ == "__main__":
    main()
