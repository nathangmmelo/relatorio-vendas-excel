import pandas as pd
import pytest
from openpyxl import load_workbook

from src.relatorio import (
    carregar_dados,
    gerar_relatorio,
    limpar_dados,
    resumo_por_produto,
    resumo_por_vendedor,
)


@pytest.fixture()
def bruto():
    return pd.DataFrame(
        {
            "data": ["2026-01-05", "2026-01-05", "2026-02-10", "2026-02-11", "data-invalida"],
            "produto": ["A", "A", "B", "B", "A"],
            "vendedor": ["ana", "ana", " BRUNO ", "Bruno", "ana"],
            "quantidade": [2, 2, 1, None, 1],
            "preco_unitario": [10.0, 10.0, 50.0, 50.0, 10.0],
        }
    )


def test_limpeza_remove_duplicata_nulos_e_data_invalida(bruto):
    limpo = limpar_dados(bruto)
    assert len(limpo) == 2


def test_limpeza_padroniza_vendedor_e_calcula_total(bruto):
    limpo = limpar_dados(bruto)
    assert set(limpo["vendedor"]) == {"Ana", "Bruno"}
    assert limpo.loc[limpo["produto"] == "A", "total"].iloc[0] == 20.0


def test_resumos(bruto):
    limpo = limpar_dados(bruto)
    por_produto = resumo_por_produto(limpo)
    assert por_produto.iloc[0]["produto"] == "B"  # 50 > 20
    assert resumo_por_vendedor(limpo)["faturamento"].sum() == 70.0


def test_colunas_ausentes_levanta_erro(tmp_path):
    arq = tmp_path / "ruim.csv"
    arq.write_text("data,produto\n2026-01-01,A\n")
    with pytest.raises(ValueError):
        carregar_dados(str(arq))


def test_gera_excel_com_abas_e_grafico(bruto, tmp_path):
    limpo = limpar_dados(bruto)
    destino = gerar_relatorio(limpo, str(tmp_path / "rel.xlsx"))
    wb = load_workbook(destino)
    assert wb.sheetnames == ["Resumo", "Por Produto", "Por Vendedor", "Mensal", "Dados Limpos"]
    assert len(wb["Por Produto"]._charts) == 1
